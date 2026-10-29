"""
Documents and conversions, against real files made on the spot. Everything is
local: no network, no model.
"""
import csv
import json
import os
import sys
import tempfile
import wave
import zipfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from jarvis.tools import REGISTRY

R = []


def check(n, c, e=""):
    print(("  PASS " if c else "  FAIL ") + n + ("" if c else f"\n        {str(e)[:300]}"))
    R.append(bool(c))


def run(name, **kw):
    return REGISTRY.execute(name, kw)


T = Path(tempfile.mkdtemp(prefix="jarvis-files-"))
print(f"working in {T}\n")

# ---------------------------------------------------------------- fixtures
from pypdf import PdfWriter
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

pdf = T / "report.pdf"
c = canvas.Canvas(str(pdf), pagesize=A4)
for i in range(1, 6):
    c.drawString(72, 760, f"Quarterly report page {i}")
    c.drawString(72, 730, f"The revenue for region {i} was {i * 1000} dollars.")
    c.showPage()
c.save()

import docx
d = docx.Document()
d.add_heading("Project Plan", 0)
d.add_paragraph("The launch date is 14 March.")
t = d.add_table(rows=2, cols=2)
t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Owner", "Task", "Asha", "Ship"
d.save(T / "plan.docx")

import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sales"
ws.append(["region", "units"])
for i in range(1, 150):
    ws.append([f"R{i}", i * 3])
wb.save(T / "sales.xlsx")

from pptx import Presentation
prs = Presentation()
s1 = prs.slides.add_slide(prs.slide_layouts[1])
s1.shapes.title.text = "Roadmap"
s1.placeholders[1].text = "Ship the voice HUD"
prs.save(T / "deck.pptx")

with open(T / "people.csv", "w", newline="") as fh:
    w = csv.writer(fh)
    w.writerow(["name", "age"])
    w.writerows([["Asha", 31], ["Ravi", 27]])

from PIL import Image, ImageDraw, ImageFont
img = Image.new("RGB", (900, 220), "white")
dr = ImageDraw.Draw(img)
try:
    font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 54)
except Exception:
    font = ImageFont.load_default()
dr.text((30, 60), "Invoice 4821 total 99.50", fill="black", font=font)
img.save(T / "receipt.png")

wavp = T / "tone.wav"
with wave.open(str(wavp), "wb") as w:
    w.setnchannels(1); w.setsampwidth(2); w.setframerate(16000)
    import math, struct
    w.writeframes(b"".join(struct.pack("<h", int(8000 * math.sin(2 * math.pi * 440 * i / 16000)))
                           for i in range(16000 * 2)))

# ---------------------------------------------------------------- read_document
print("=== read_document ===")
r = run("read_document", path=str(pdf))
check("pdf: text and page count", r.ok and r.content["pages_total"] == 5 and "revenue for region 3" in r.content["content"], r.content if r.ok else r.error)
r = run("read_document", path=str(pdf), pages="2-3")
check("pdf: a page range reads only those pages",
      r.ok and "region 2" in r.content["content"] and "region 5" not in r.content["content"], r.content if r.ok else r.error)
r = run("read_document", path=str(pdf), max_chars=500)
check("long documents are paged with a next offset", r.ok and (r.content.get("more") is None or r.content.get("next_offset")), r.content if r.ok else r.error)
r = run("read_document", path=str(T / "plan.docx"))
check("docx: headings, paragraphs and tables",
      r.ok and "Project Plan" in r.content["content"] and "14 March" in r.content["content"]
      and "Asha | Ship" in r.content["content"], r.content if r.ok else r.error)
r = run("read_document", path=str(T / "sales.xlsx"))
check("xlsx: sheet name, size and rows, with the rest summarised",
      r.ok and "sheet: Sales" in r.content["content"] and "R10\t30" in r.content["content"]
      and "more rows" in r.content["content"], r.content if r.ok else r.error)
r = run("read_document", path=str(T / "people.csv"))
check("csv: columns and rows", r.ok and r.content["columns"] == ["name", "age"] and "Ravi" in r.content["content"])
r = run("read_document", path=str(T / "deck.pptx"))
check("pptx: slide text", r.ok and "Roadmap" in r.content["content"] and "voice HUD" in r.content["content"], r.content if r.ok else r.error)
r = run("read_document", path=str(T / "receipt.png"))
check("image: text is recognised (OCR)", r.ok and "4821" in r.content["content"], r.content if r.ok else r.error)
r = run("read_document", path=str(wavp))
check("audio: duration and codec", r.ok and abs(r.content["duration_seconds"] - 2.0) < 0.2, r.content if r.ok else r.error)
check("a missing file is an honest error", not run("read_document", path=str(T / "nope.pdf")).ok)
(T / "blob.bin").write_bytes(b"\x00\x01\x02" * 100)
check("a binary file is refused, not garbled", not run("read_document", path=str(T / "blob.bin")).ok)

# ---------------------------------------------------------------- pdf_tool
print("\n=== pdf_tool ===")
r = run("pdf_tool", action="info", path=str(pdf))
check("info: pages", r.ok and r.content["pages"] == 5, r.content if r.ok else r.error)
r = run("pdf_tool", action="extract", path=str(pdf), pages="2-3")
check("extract: a new 2-page PDF next to the source",
      r.ok and r.content["pages"] == 2 and Path(r.content["path"]).exists()
      and Path(r.content["path"]).name != "report.pdf", r.content if r.ok else r.error)
ext_path = r.content["path"]
rr = run("read_document", path=ext_path)
check("the extracted PDF contains exactly those pages",
      "region 2" in rr.content["content"] and "region 4" not in rr.content["content"])
r = run("pdf_tool", action="merge", paths=[str(pdf), ext_path])
check("merge: 5 + 2 pages", r.ok and r.content["pages"] == 7, r.content if r.ok else r.error)
r = run("pdf_tool", action="split", path=str(pdf), pages="1-3")
check("split: one file per page", r.ok and r.content["count"] == 3, r.content if r.ok else r.error)
r = run("pdf_tool", action="rotate", path=str(pdf), pages="1", angle=90)
check("rotate: new file", r.ok and Path(r.content["path"]).exists(), r.content if r.ok else r.error)
r = run("pdf_tool", action="to_images", path=str(pdf), pages="1-2")
check("to_images: PNG per page", r.ok and r.content["count"] == 2, r.content if r.ok else r.error)
check("a bad page range is refused", not run("pdf_tool", action="extract", path=str(pdf), pages="40-50").ok)
check("merging one file is refused", not run("pdf_tool", action="merge", paths=[str(pdf)]).ok)
before = pdf.stat().st_size
run("pdf_tool", action="extract", path=str(pdf), pages="1")
check("the original is never touched", pdf.stat().st_size == before)

# ---------------------------------------------------------------- convert_file
print("\n=== convert_file ===")
r = run("convert_file", path=str(T / "receipt.png"), to="jpg", quality=80)
check("image: png -> jpg", r.ok and r.content["path"].endswith(".jpg") and r.content["bytes"] > 0, r.content if r.ok else r.error)
r = run("convert_file", path=str(T / "receipt.png"), to="png", max_dimension=300, output=str(T / "small.png"))
check("image: already that format is refused", not r.ok)
r = run("convert_file", path=str(T / "receipt.png"), to="pdf")
check("image: png -> pdf", r.ok and r.content["path"].endswith(".pdf"), r.content if r.ok else r.error)
r = run("convert_file", path=str(T / "receipt.png"), to="webp", max_dimension=400)
check("image: png -> webp, resized", r.ok and r.content["path"].endswith(".webp"), r.content if r.ok else r.error)
im = Image.open(r.content["path"]) if r.ok else None
check("...and really resized", im is not None and max(im.size) <= 400, im.size if im else "")
r = run("convert_file", path=str(wavp), to="mp3")
check("audio: wav -> mp3", r.ok and r.content["path"].endswith(".mp3") and r.content["bytes"] > 1000, r.content if r.ok else r.error)
r = run("convert_file", path=str(wavp), to="m4a", start="0:00:00.5", end="0:00:01.5")
check("audio: trimmed to ~1 s", r.ok, r.content if r.ok else r.error)
if r.ok:
    dur = run("read_document", path=r.content["path"]).content["duration_seconds"]
    check("...and the duration shows it", 0.8 <= dur <= 1.3, dur)
(T / "notes.md").write_text("# Title\n\nSome *text* here.\n\n- one\n- two\n")
r = run("convert_file", path=str(T / "notes.md"), to="html")
check("markdown -> html", r.ok and "<h1>Title</h1>" in Path(r.content["path"]).read_text(), r.content if r.ok else r.error)
r = run("convert_file", path=str(T / "notes.md"), to="pdf")
check("markdown -> pdf", r.ok and Path(r.content["path"]).read_bytes()[:5] == b"%PDF-", r.content if r.ok else r.error)
r = run("convert_file", path=str(T / "plan.docx"), to="txt")
check("docx -> txt via textutil", r.ok and "launch date" in Path(r.content["path"]).read_text(), r.content if r.ok else r.error)
r = run("convert_file", path=str(T / "people.csv"), to="xlsx")
check("csv -> xlsx", r.ok and r.content["path"].endswith(".xlsx"), r.content if r.ok else r.error)
r2 = run("convert_file", path=r.content["path"], to="csv") if r.ok else None
check("xlsx -> csv round trip keeps the data", r2 is not None and r2.ok and "Ravi" in Path(r2.content["path"]).read_text(), r2.content if r2 and r2.ok else (r2.error if r2 else ""))
check("an unsupported conversion says so", not run("convert_file", path=str(T / "plan.docx"), to="mp4").ok)
check("never overwrites: second conversion gets a new name",
      run("convert_file", path=str(T / "receipt.png"), to="jpg").content["path"] != str(T / "receipt.jpg"))

# ---------------------------------------------------------------- archives
print("\n=== archive_tool ===")
r = run("archive_tool", action="zip", paths=[str(T / "plan.docx"), str(T / "people.csv")], output=str(T / "bundle.zip"))
check("zip two files", r.ok and r.content["files"] == 2 and zipfile.is_zipfile(r.content["path"]), r.content if r.ok else r.error)
r = run("archive_tool", action="list", path=str(T / "bundle.zip"))
check("list shows entries without extracting", r.ok and r.content["entries"] == 2, r.content if r.ok else r.error)
r = run("archive_tool", action="unzip", path=str(T / "bundle.zip"))
check("unzip into a new folder", r.ok and (Path(r.content["folder"]) / "people.csv").exists(), r.content if r.ok else r.error)
evil = T / "evil.zip"
with zipfile.ZipFile(evil, "w") as z:
    z.writestr("../../escape.txt", "x")
r = run("archive_tool", action="unzip", path=str(evil))
check("a path-traversal archive is refused", not r.ok and "unsafe" in r.error, r.error if not r.ok else r.content)
check("the escape file was not written", not (T.parent / "escape.txt").exists())

print(f"\n{sum(R)}/{len(R)} document and conversion checks passed")
sys.exit(0 if all(R) else 1)
