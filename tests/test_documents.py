"""
make_document: designed PDFs with diagrams, charts and pictures, built for real
with headless Chrome. Uses a local picture so no image search is needed; the
first run downloads Mermaid once into ~/.jarvis/cache.
"""
import os
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from jarvis.tools import REGISTRY
from jarvis.tools import documents_pro as dp

R = []


def check(n, c, e=""):
    print(("  PASS " if c else "  FAIL ") + n + ("" if c else f"\n        {str(e)[:400]}"))
    R.append(bool(c))


def run(name, **kw):
    return REGISTRY.execute(name, kw)


T = Path(tempfile.mkdtemp(prefix="jarvis-docs-"))
print(f"working in {T}\n")

from PIL import Image, ImageDraw

pic = T / "site.png"
im = Image.new("RGB", (1600, 900), (40, 90, 160))
ImageDraw.Draw(im).rectangle((200, 300, 1400, 700), fill=(240, 200, 60))
im.save(pic)

BODY = f"""# Pilot Plan

## Summary
A short plan with **every** kind of figure.

> [!TIP]
> Start small.

![Site photo]({pic})

## Flow
```mermaid
%% caption: How a request moves
flowchart LR
  A[Request] --> B{{Approved?}}
  B -->|yes| C[Build]
  B -->|no| D[Revise]
```

## Numbers
```chart
{{"kind": "bar", "title": "Cost", "categories": ["Q1", "Q2"],
  "series": [{{"name": "Plan", "values": [3, 5]}}]}}
```

| Item | Cost |
|---|---|
| Panels | 68 |

## Architecture
```mermaid
%% caption: Parts
flowchart TB
  subgraph Site
    P[Panels] --> I[Inverter]
  end
  I --> G[Gateway]
```
"""

# ---------------------------------------------------------------- unit: ordering
print("=== figure numbering ===")
b = dp.Build(T / "unit")
body, heads = dp.markdown_to_body(BODY, b)
labels = [int(x) for x in __import__("re").findall(r"fig-label'>Figure (\d+)<", body)]
check("figures are numbered 1..n in reading order", labels == list(range(1, len(labels) + 1)) and len(labels) == 4, labels)
first_fig = body.find("fig-label'>Figure 1<")
check("the photo (first in the text) is Figure 1", 0 <= first_fig < body.find("class='mermaid'"), body[:300])
check("top-to-bottom diagrams are marked tall", "diagram tall" in body)
check("sections are collected for the contents", [h[1] for h in heads if h[0] == 2][:2] == ["Summary", "Flow"], heads)

# ---------------------------------------------------------------- full build
print("\n=== full PDF build ===")
r = run("make_document", path=str(T / "Pilot Plan.pdf"), title="Pilot Plan", subtitle="A test",
        kicker="Plan", markdown=BODY)
check("builds", r.ok, r.error)
if r.ok:
    c = r.content
    check("pdf exists and is verified", c["verified"] and Path(c["path"]).exists() and c["bytes"] > 20_000, c)
    check("counts the figures", c["diagrams"] == 2 and c["charts"] == 1 and c["photos"] == 1, c)
    check("has a cover, contents and body pages", c.get("pages", 0) >= 3, c)
    check("a preview sheet was made", Path(c.get("preview", "/nonexistent")).exists(), c)
    from pypdf import PdfReader
    text = "".join(p.extract_text() or "" for p in PdfReader(c["path"]).pages)
    check("the title is on the cover", "Pilot Plan" in text)
    flat = "".join(text.split())                 # extraction can split a word with a space
    check("diagram labels are real text in the PDF", "Request" in flat and "Inverter" in flat, text[:400])
    check("pages are numbered", "1 / " in text or "1/" in text, text[-300:])
    # once on the cover, then only in each later page's footer
    # (whitespace-insensitive: PDF text extraction may wrap the cover title onto two lines)
    check("the title heading is not repeated in the body", flat.count("PilotPlan") == c["pages"], (flat.count("PilotPlan"), c["pages"]))
    check("Mermaid is cached locally for later builds", dp.MERMAID_LOCAL.exists())

    r2 = run("make_document", path=str(T / "Pilot Plan.pdf"), title="Pilot Plan", markdown="## A\nx\n")
    check("never overwrites an existing file", r2.ok and r2.content["path"] != c["path"], r2.content if r2.ok else r2.error)

# ---------------------------------------------------------------- broken diagram
print("\n=== broken diagram ===")
bad = "## X\n```mermaid\nflowchart LR\n  A[Start --> --> ((B\n  C -->\n```\n"
r = run("make_document", path=str(T / "Broken.pdf"), title="Broken", markdown=bad)
check("a diagram Mermaid cannot draw is refused", not r.ok and "diagram" in (r.error or ""), r.content if r.ok else r.error)
check("the error carries the source so it can be fixed", not r.ok and "flowchart LR" in (r.error or ""), r.error)
check("no PDF is left behind", not (T / "Broken.pdf").exists())

# ---------------------------------------------------------------- html + themes
print("\n=== html and themes ===")
r = run("make_document", path=str(T / "Page.html"), title="Page", theme="warm", markdown=BODY)
check("html output works", r.ok and r.content["path"].endswith(".html"), r.error)
if r.ok:
    html_text = Path(r.content["path"]).read_text()
    check("a shared html file loads Mermaid from the web, not this Mac", "file://" not in html_text and "cdn.jsdelivr.net" in html_text)
    check("its picture travels with it", (T / "Page_files").is_dir() and any((T / "Page_files").iterdir()))
r = run("make_document", path=str(T / "Classic.pdf"), title="Classic", theme="classic", cover=False,
        markdown="## One\ntext\n## Two\ntext\n")
check("classic theme without a cover", r.ok and r.content.get("pages", 0) >= 1, r.error)
r = run("make_document", path=str(T / "odd.xyz"), title="Odd", markdown="## A\nb\n")
check("an unsupported extension becomes a PDF", r.ok and r.content["path"].endswith(".pdf"), r.error)

print("\n=== Word (.docx) ===")
import re as _re
import zipfile
from docx import Document as _Doc
r = run("make_document", path=str(T / "Pilot Plan.docx"), title="Pilot Plan", subtitle="A test",
        author="JARVIS", markdown=BODY)
check("a .docx path makes a Word document", r.ok and r.content["path"].endswith(".docx"),
      r.error if not r.ok else r.content)
if r.ok:
    d = _Doc(r.content["path"])
    styles = [p.style.name for p in d.paragraphs]
    xml = zipfile.ZipFile(r.content["path"]).read("word/document.xml").decode()
    check("sections are real Word headings (navigation pane, contents)",
          styles.count("Heading 1") >= 3, styles[:20])
    check("tables, charts and diagrams come through", len(d.tables) >= 1 and "<pic:pic" in xml
          and len(d.inline_shapes) >= r.content["charts"] + r.content["diagrams"],
          (len(d.tables), len(d.inline_shapes), r.content))
    check("figures are captioned", "Caption" in styles)
    foot = zipfile.ZipFile(r.content["path"]).read("word/footer1.xml").decode() \
        if "word/footer1.xml" in zipfile.ZipFile(r.content["path"]).namelist() else ""
    check("pages are numbered", "PAGE" in foot and "NUMPAGES" in foot)
    order_ok = all(m.group(1).index("tblW") < m.group(1).index("tblLook")
                   for m in _re.finditer(r"<w:tblPr>(.*?)</w:tblPr>", xml))
    check("table properties are in schema order (Word opens it without 'repair')", order_ok)
    check("a preview of the first page comes back", r.content.get("preview") and
          Path(r.content["preview"]).exists(), r.content)
r = run("make_document", path=str(T / "Lists.docx"), title="Lists", cover=False,
        markdown="## Steps\n- read\n- write\n    - nested\n\n1. one\n2. two\n")
ls = [p.style.name for p in _Doc(r.content["path"]).paragraphs] if r.ok else []
check("lists are real Word lists, nesting included",
      ls.count("List Bullet") == 2 and "List Bullet 2" in ls and ls.count("List Number") == 2, ls)
r = run("make_document", path=str(T / "Letter.word"), title="Letter", markdown="Dear team,\n\nThanks.", cover=False)
check("'word' as the extension also makes .docx", r.ok and r.content["path"].endswith(".docx"), r.error)

print(f"\n{sum(R)}/{len(R)} document checks passed")
sys.exit(0 if all(R) else 1)
