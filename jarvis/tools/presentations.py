"""
Presentations: build and edit real .pptx files (open in PowerPoint or Keynote).

The model supplies *content* - an outline of slides - and this module owns the
*design*: consistent type scale, margins, colour themes, text sized to fit, and
layouts for the things decks actually contain (bullets, two columns, images,
quotes, tables, charts, big numbers). That split is what makes the output look
designed instead of like a text dump, and keeps a 12-slide deck to a single
tool call.

Nothing is overwritten: `make_presentation` refuses an existing path (it picks
a new name), and `edit_presentation` writes a copy unless told otherwise, in
which case the original is snapshotted so the edit can be undone.
"""
from __future__ import annotations

import logging
import os
import shutil
import subprocess
import time
from pathlib import Path

from ..safety import journal
from ..safety.policy import Risk, Verdict, classify_path_write
from .fs import _p
from .registry import REGISTRY, ToolResult

log = logging.getLogger("jarvis.presentations")

W_IN, H_IN = 13.333, 7.5          # 16:9

THEMES: dict[str, dict] = {
    # bg, text, muted, accent, accent2, font (head), font (body)
    "modern": dict(bg="FFFFFF", text="1D1D1F", muted="6E6E73", accent="0A84FF",
                   accent2="5E5CE6", panel="F5F5F7", head="Helvetica Neue", body="Helvetica Neue"),
    "dark": dict(bg="111114", text="F5F5F7", muted="A1A1A6", accent="64D2FF",
                 accent2="BF5AF2", panel="1C1C21", head="Helvetica Neue", body="Helvetica Neue"),
    "minimal": dict(bg="FAFAF7", text="222222", muted="777777", accent="222222",
                    accent2="C0392B", panel="EFEFEA", head="Georgia", body="Helvetica Neue"),
    "bold": dict(bg="FFFFFF", text="101010", muted="555555", accent="FF3B30",
                 accent2="FF9500", panel="FFF1EE", head="Avenir Next", body="Avenir Next"),
    "warm": dict(bg="FFF9F2", text="3A2A1F", muted="8C7B6B", accent="D9622B",
                 accent2="2E7D6B", panel="FBEBD9", head="Georgia", body="Avenir Next"),
    "forest": dict(bg="F4F8F5", text="12261C", muted="5C7266", accent="1F7A4D",
                   accent2="C9A227", panel="E3EEE7", head="Avenir Next", body="Avenir Next"),
}
DEFAULT_THEME = "modern"
LAYOUTS = ("title", "bullets", "two_column", "image", "image_left", "image_right",
           "quote", "section", "table", "chart", "big_number", "closing")


def _rgb(hexs: str):
    from pptx.dml.color import RGBColor
    return RGBColor.from_string(hexs.upper())


def _theme(name: str) -> dict:
    return THEMES.get((name or DEFAULT_THEME).lower(), THEMES[DEFAULT_THEME])


# ------------------------------------------------------------------ drawing helpers

def _rect(slide, x, y, w, h, fill, line=False):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    if not line:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def _text(slide, x, y, w, h, text, size, color, font, bold=False, italic=False,
          align="left", anchor="top", line_spacing=1.1):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE,
                          "bottom": MSO_ANCHOR.BOTTOM}[anchor]
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
                       "right": PP_ALIGN.RIGHT}[align]
        p.line_spacing = line_spacing
        r = p.add_run()
        r.text = str(line)
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.name = font
        f.color.rgb = _rgb(color)
    return tb


def _fit(chars: int, base: int, lo: int = 15) -> int:
    """Shrink the type as the text grows so it stays inside its box."""
    if chars > 900:
        return max(lo, base - 10)
    if chars > 650:
        return max(lo, base - 7)
    if chars > 420:
        return max(lo, base - 4)
    if chars > 260:
        return max(lo, base - 2)
    return base


def _bullets(slide, x, y, w, h, items, th, base=24, color=None):
    """A bulleted list; items are strings or {text, sub: [...]}."""
    from pptx.util import Inches, Pt
    flat: list[tuple[str, int]] = []
    for it in items or []:
        if isinstance(it, dict):
            flat.append((str(it.get("text", "")), 0))
            for s in it.get("sub", []) or []:
                flat.append((str(s), 1))
        else:
            flat.append((str(it), 0))
    chars = sum(len(t) for t, _ in flat)
    size = _fit(chars + 18 * len(flat), base)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (t, lvl) in enumerate(flat):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(size * (0.45 if lvl == 0 else 0.25))
        p.line_spacing = 1.08
        bullet = "•  " if lvl == 0 else "–  "
        r0 = p.add_run()
        r0.text = ("      " if lvl else "") + bullet
        r0.font.size = Pt(size - (2 if lvl else 0))
        r0.font.name = th["body"]
        r0.font.color.rgb = _rgb(th["accent"])
        r0.font.bold = True
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size - (3 if lvl else 0))
        r.font.name = th["body"]
        r.font.color.rgb = _rgb(color or (th["text"] if lvl == 0 else th["muted"]))
    return tb


def _title_bar(slide, title, th, y=0.55):
    _rect(slide, 0.8, y + 0.07, 0.09, 0.78, th["accent"])
    size = 36 if len(title) <= 42 else 30 if len(title) <= 64 else 26
    _text(slide, 1.05, y, 11.4, 0.95, title, size, th["text"], th["head"],
          bold=True, anchor="middle")


def _picture(slide, path, x, y, w, h):
    """Place an image inside a box, keeping its proportions, centred."""
    from PIL import Image
    from pptx.util import Inches
    try:
        iw, ih = Image.open(path).size
    except Exception:
        return None
    scale = min(w / iw, h / ih)
    pw, ph = iw * scale, ih * scale
    return slide.shapes.add_picture(str(path), Inches(x + (w - pw) / 2),
                                    Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))


def _footer(slide, n, total, th, deck_title=""):
    if deck_title:
        _text(slide, 0.8, 7.0, 8, 0.3, deck_title, 10, th["muted"], th["body"])
    _text(slide, 11.6, 7.0, 0.95, 0.3, f"{n}", 10, th["muted"], th["body"], align="right")


# ------------------------------------------------------------------ layouts

def _bg(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(color)


def _notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = str(text)


def render_slide(prs, spec: dict, th: dict, number: int, total: int, deck_title: str,
                 warnings: list[str]):
    from pptx.util import Inches
    layout = str(spec.get("layout") or ("bullets" if spec.get("bullets") else "title"))
    if layout not in LAYOUTS:
        warnings.append(f"slide {number}: unknown layout {layout!r}, used bullets")
        layout = "bullets"
    s = prs.slides.add_slide(prs.slide_layouts[6])        # blank
    title = str(spec.get("title", "")).strip()
    sub = str(spec.get("subtitle", "")).strip()

    if layout == "title":
        _bg(s, th["bg"])
        _rect(s, 0, 0, 0.35, H_IN, th["accent"])
        _text(s, 1.1, 2.3, 11.2, 1.9, title, 54 if len(title) < 40 else 44, th["text"],
              th["head"], bold=True, anchor="bottom")
        _rect(s, 1.15, 4.35, 1.6, 0.07, th["accent2"])
        if sub:
            _text(s, 1.1, 4.6, 11, 1.2, sub, 24, th["muted"], th["body"])
    elif layout == "section":
        _bg(s, th["accent"])
        _text(s, 1.1, 2.6, 11.2, 1.8, title, 50, "FFFFFF", th["head"], bold=True,
              anchor="middle")
        if sub:
            _text(s, 1.1, 4.4, 11, 1.0, sub, 22, "FFFFFF", th["body"])
    elif layout == "closing":
        _bg(s, th["bg"])
        _text(s, 1.1, 2.6, 11.2, 1.6, title or "Thank you", 54, th["text"], th["head"],
              bold=True, align="center", anchor="middle")
        _rect(s, 5.9, 4.3, 1.5, 0.07, th["accent"])
        if sub:
            _text(s, 1.1, 4.6, 11.2, 1.0, sub, 22, th["muted"], th["body"], align="center")
    elif layout == "quote":
        _bg(s, th["panel"])
        _text(s, 1.3, 1.2, 1.5, 1.5, "“", 120, th["accent"], th["head"], bold=True)
        q = str(spec.get("quote") or title)
        _text(s, 1.6, 2.2, 10.2, 3.2, q, _fit(len(q), 36, 24), th["text"], th["head"],
              italic=True, anchor="middle", line_spacing=1.2)
        if spec.get("author"):
            _text(s, 1.6, 5.6, 10, 0.6, "— " + str(spec["author"]), 20, th["muted"], th["body"])
    elif layout == "big_number":
        _bg(s, th["bg"])
        num = str(spec.get("number", title))
        _text(s, 0.8, 1.6, 11.7, 2.8, num, 140 if len(num) < 7 else 96, th["accent"],
              th["head"], bold=True, align="center", anchor="middle")
        lab = str(spec.get("label") or sub)
        if lab:
            _text(s, 1.5, 4.5, 10.3, 1.2, lab, 28, th["text"], th["body"], align="center")
        if title and spec.get("number"):
            _text(s, 1.5, 0.7, 10.3, 0.8, title, 22, th["muted"], th["body"], align="center")
    else:
        _bg(s, th["bg"])
        if title:
            _title_bar(s, title, th)
        top = 1.9
        if layout == "bullets":
            _bullets(s, 1.05, top, 11.3, 4.9, spec.get("bullets"), th, 28)
        elif layout == "two_column":
            rows_n = max(len(spec.get("left") or []), len(spec.get("right") or []))
            panel_h = max(3.0, min(4.8, 1.5 + 0.8 * rows_n))
            for col, key, tkey in ((0, "left", "left_title"), (1, "right", "right_title")):
                x = 0.9 + col * 6.0
                _rect(s, x, top, 5.6, panel_h, th["panel"])
                y = top + 0.2
                if spec.get(tkey):
                    _text(s, x + 0.25, y, 5.1, 0.6, str(spec[tkey]), 22, th["accent"],
                          th["head"], bold=True)
                    y += 0.75
                _bullets(s, x + 0.25, y, 5.1, panel_h - (y - top) - 0.2, spec.get(key), th, 23)
        elif layout in ("image", "image_left", "image_right"):
            img = spec.get("image")
            has_text = layout != "image" and (spec.get("bullets") or spec.get("caption"))
            if img and Path(str(img)).expanduser().exists():
                ip = str(Path(str(img)).expanduser())
                if layout == "image" or not has_text:
                    _picture(s, ip, 0.9, top, 11.5, 4.5 if spec.get("caption") else 4.9)
                else:
                    left = layout == "image_left"
                    _picture(s, ip, 0.9 if left else 6.9, top, 5.6, 4.8)
                    _bullets(s, 6.9 if left else 0.9, top, 5.6, 4.8,
                             spec.get("bullets") or [spec.get("caption", "")], th, 22)
            else:
                warnings.append(f"slide {number}: image not found: {img!r}")
                _bullets(s, 1.05, top, 11.3, 4.9, spec.get("bullets") or [spec.get("caption", "")],
                         th, 24)
            if layout == "image" and spec.get("caption"):
                _text(s, 0.9, 6.5, 11.5, 0.5, str(spec["caption"]), 16, th["muted"],
                      th["body"], align="center", italic=True)
        elif layout == "table":
            rows = spec.get("table") or []
            if rows:
                nr, nc = len(rows), max(len(r) for r in rows)
                row_h = 0.8 if nr <= 5 else 0.62 if nr <= 8 else 0.5
                gs = s.shapes.add_table(nr, nc, Inches(0.9), Inches(top), Inches(11.5),
                                        Inches(min(4.9, row_h * nr)))
                tbl = gs.table
                from pptx.util import Pt
                for ri, row in enumerate(rows):
                    for ci in range(nc):
                        cell = tbl.cell(ri, ci)
                        cell.text = str(row[ci]) if ci < len(row) else ""
                        para = cell.text_frame.paragraphs[0]
                        para.runs[0].font.size = Pt(22 if nr <= 5 else 18 if nr <= 8 else 14)
                        para.runs[0].font.name = th["body"]
                        cell.fill.solid()
                        if ri == 0:
                            cell.fill.fore_color.rgb = _rgb(th["accent"])
                            para.runs[0].font.color.rgb = _rgb("FFFFFF")
                            para.runs[0].font.bold = True
                        else:
                            cell.fill.fore_color.rgb = _rgb(th["panel"] if ri % 2 else th["bg"])
                            para.runs[0].font.color.rgb = _rgb(th["text"])
            else:
                warnings.append(f"slide {number}: table has no rows")
        elif layout == "chart":
            _chart(s, spec.get("chart") or {}, th, top, warnings, number)

    _notes(s, spec.get("notes"))
    if layout not in ("title", "section", "closing"):
        _footer(s, number, total, th, deck_title)
    if layout == "bullets" and len(spec.get("bullets") or []) > 7:
        warnings.append(f"slide {number}: {len(spec['bullets'])} bullets is a lot - consider splitting")
    return s


def _chart(slide, spec: dict, th: dict, top: float, warnings: list, number: int):
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Inches, Pt
    cats = [str(c) for c in spec.get("categories", [])]
    series = spec.get("series") or []
    if not cats or not series:
        warnings.append(f"slide {number}: chart needs categories and series")
        return
    kind = str(spec.get("type", "column")).lower()
    ctype = {"bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
             "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE,
             "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA,
             "stacked": XL_CHART_TYPE.COLUMN_STACKED}.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED)
    cd = CategoryChartData()
    cd.categories = cats
    for sr in series:
        cd.add_series(str(sr.get("name", "")), [float(v) for v in sr.get("values", [])])
    gf = slide.shapes.add_chart(ctype, Inches(0.9), Inches(top), Inches(11.5), Inches(4.8), cd)
    ch = gf.chart
    ch.font.size = Pt(14)
    ch.font.name = th["body"]
    ch.font.color.rgb = _rgb(th["text"])
    ch.has_legend = len(series) > 1 or kind in ("pie", "doughnut")
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
    palette = [th["accent"], th["accent2"], "34C759", "FF9500", "AF52DE", "FF3B30"]
    try:
        if kind in ("pie", "doughnut"):
            for i, pt in enumerate(ch.plots[0].series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
        else:
            for i, sr in enumerate(ch.plots[0].series):
                sr.format.fill.solid()
                sr.format.fill.fore_color.rgb = _rgb(palette[i % len(palette)])
                if kind == "line":
                    sr.format.line.color.rgb = _rgb(palette[i % len(palette)])
    except Exception:
        pass
    if spec.get("title"):
        ch.has_title = True
        ch.chart_title.text_frame.text = str(spec["title"])


# ------------------------------------------------------------------ tools

def _assess_out(args: dict) -> Verdict:
    target = args.get("output") or args.get("path") or ""
    if target:
        v = classify_path_write(str(target))
        if v.risk >= Risk.FORBIDDEN:
            return v
    return Verdict(Risk.LOW, "creates a presentation file", category="write_new")


def _unique(p: Path) -> Path:
    if not p.exists():
        return p
    for i in range(2, 1000):
        c = p.with_name(f"{p.stem} ({i}){p.suffix}")
        if not c.exists():
            return c
    return p.with_name(f"{p.stem}-{int(time.time())}{p.suffix}")


def thumbnail(path: Path, size: int = 1000) -> str | None:
    """A PNG of the first slide via Quick Look, for a quick visual check."""
    try:
        out = Path(os.path.expanduser("~/.jarvis/data/previews"))
        out.mkdir(parents=True, exist_ok=True)
        r = subprocess.run(["/usr/bin/qlmanage", "-t", "-s", str(size), "-o", str(out), str(path)],
                           capture_output=True, text=True, timeout=40)
        png = out / (path.name + ".png")
        return str(png) if png.exists() else None
    except Exception:
        return None


@REGISTRY.tool(
    "make_presentation",
    "Create a designed PowerPoint (.pptx) deck - opens in PowerPoint or Keynote. "
    "Give the CONTENT as an ordered list of slides; the design is handled. Each "
    "slide: layout (title | bullets | two_column | image | image_left | "
    "image_right | quote | section | table | chart | big_number | closing), "
    "title, and the fields that layout uses: bullets (strings, or {text, sub}), "
    "left/right (+left_title/right_title), image (path) + caption, quote + "
    "author, table (rows, first row is the header), chart {type: column|bar|"
    "line|pie, categories, series:[{name, values}]}, number + label, subtitle, "
    "notes (speaker notes). Good decks: one idea per slide, 3-5 bullets of a "
    "few words each, a title slide first and a closing slide last, speaker "
    "notes for the detail. theme: modern, dark, minimal, bold, warm, forest.",
    {"type": "object", "properties": {
        "path": {"type": "string", "description": "Where to save, e.g. ~/Desktop/Pitch.pptx"},
        "slides": {"type": "array", "items": {"type": "object"}},
        "theme": {"type": "string", "enum": list(THEMES)},
        "title": {"type": "string", "description": "Deck title (footer, file properties)"}},
     "required": ["path", "slides"]},
    risk=Risk.LOW, category="write_new", parallel_safe=False, slow=True, assess=_assess_out)
def make_presentation(path: str, slides: list, theme: str = DEFAULT_THEME, title: str = ""):
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        return ToolResult(False, error="python-pptx is not installed")
    if not isinstance(slides, list) or not slides:
        return ToolResult(False, error="give 'slides' as a non-empty list")
    if len(slides) > 60:
        return ToolResult(False, error="more than 60 slides; split the deck")
    f = _p(path)
    if f.suffix.lower() != ".pptx":
        f = f.with_suffix(".pptx")
    f = _unique(f)
    f.parent.mkdir(parents=True, exist_ok=True)
    th = _theme(theme)
    tname = (theme or DEFAULT_THEME).lower() if (theme or "").lower() in THEMES else DEFAULT_THEME
    warnings: list[str] = []
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W_IN), Inches(H_IN)
    deck_title = title or next((str(s.get("title", "")) for s in slides
                                if isinstance(s, dict) and s.get("title")), "")
    try:
        for i, spec in enumerate(slides, 1):
            if not isinstance(spec, dict):
                warnings.append(f"slide {i}: not an object, skipped")
                continue
            render_slide(prs, spec, th, i, len(slides), deck_title, warnings)
        prs.core_properties.title = deck_title
        prs.core_properties.keywords = f"jarvis-theme:{tname}"
        prs.core_properties.author = "JARVIS"
        prs.save(str(f))
    except Exception as e:                        # noqa: BLE001
        return ToolResult(False, error=f"could not build the deck: {e}"[:300])
    if not f.exists() or f.stat().st_size == 0:
        return ToolResult(False, error="the file was not created")
    try:
        n = len(Presentation(str(f)).slides)
    except Exception as e:
        return ToolResult(False, error=f"the saved deck does not open: {e}"[:200])
    journal.record("create", str(f), snapshot=False, session=REGISTRY.session_id)
    content = {"path": str(f), "slides": n, "theme": tname, "bytes": f.stat().st_size,
               "verified": n == len([s for s in slides if isinstance(s, dict)])}
    if warnings:
        content["warnings"] = warnings
    content["note"] = ("offer to open it with open_file; edit later with "
                       "edit_presentation")
    return ToolResult(True, content=content,
                      display=f"created {f.name} ({n} slides, {tname} theme)")


def _text_frames(prs):
    for slide in prs.slides:
        for shp in slide.shapes:
            if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
                yield slide, shp.text_frame


@REGISTRY.tool(
    "edit_presentation",
    "Change an existing .pptx. ops is a list, applied in order: "
    "{op:'add_slide', after:N (0=first), slide:{...same fields as make_presentation}}, "
    "{op:'delete_slide', index:N}, {op:'move_slide', from:N, to:M}, "
    "{op:'replace_text', find, replace}, {op:'set_notes', index:N, notes}, "
    "{op:'set_title', index:N, title}. Slide numbers start at 1. Writes a copy "
    "next to the original unless in_place=true (the original is then "
    "snapshotted, so it can be undone).",
    {"type": "object", "properties": {
        "path": {"type": "string"},
        "ops": {"type": "array", "items": {"type": "object"}},
        "in_place": {"type": "boolean"},
        "output": {"type": "string"}},
     "required": ["path", "ops"]},
    risk=Risk.MEDIUM, category="write_new", parallel_safe=False, slow=True,
    assess=lambda a: classify_path_write(str(a.get("path", ""))) if a.get("in_place")
    else _assess_out(a))
def edit_presentation(path: str, ops: list, in_place: bool = False, output: str = ""):
    try:
        from pptx import Presentation
    except ImportError:
        return ToolResult(False, error="python-pptx is not installed")
    src = _p(path)
    if not src.exists():
        return ToolResult(False, error=f"File not found: {src}")
    try:
        prs = Presentation(str(src))
    except Exception as e:
        return ToolResult(False, error=f"cannot open the deck: {e}"[:200])
    kw = prs.core_properties.keywords or ""
    th = _theme(kw.split("jarvis-theme:")[1].split()[0] if "jarvis-theme:" in kw else DEFAULT_THEME)
    warnings: list[str] = []
    applied: list[str] = []

    def n_slides():
        return len(prs.slides)

    def move(old: int, new: int):
        lst = prs.slides._sldIdLst
        items = list(lst)
        el = items[old]
        lst.remove(el)
        lst.insert(new, el)

    for op in ops or []:
        if not isinstance(op, dict):
            continue
        kind = str(op.get("op", ""))
        try:
            if kind == "add_slide":
                after = int(op.get("after", n_slides()))
                before = n_slides()
                render_slide(prs, op.get("slide") or {}, th, after + 1, before + 1,
                             prs.core_properties.title or "", warnings)
                move(n_slides() - 1, max(0, min(after, before)))
                applied.append(f"added a slide after {after}")
            elif kind == "delete_slide":
                i = int(op["index"]) - 1
                if not 0 <= i < n_slides():
                    raise IndexError(f"no slide {i + 1}")
                lst = prs.slides._sldIdLst
                prs.part.drop_rel(list(lst)[i].rId)
                lst.remove(list(lst)[i])
                applied.append(f"deleted slide {i + 1}")
            elif kind == "move_slide":
                a, b = int(op["from"]) - 1, int(op["to"]) - 1
                if not (0 <= a < n_slides() and 0 <= b < n_slides()):
                    raise IndexError("slide number out of range")
                move(a, b)
                applied.append(f"moved slide {a + 1} to {b + 1}")
            elif kind == "replace_text":
                find, rep = str(op["find"]), str(op.get("replace", ""))
                hits = 0
                for _, tf in _text_frames(prs):
                    for para in tf.paragraphs:
                        for run in para.runs:
                            if find in run.text:
                                run.text = run.text.replace(find, rep)
                                hits += 1
                applied.append(f"replaced {find!r} in {hits} place(s)")
                if not hits:
                    warnings.append(f"{find!r} was not found")
            elif kind == "set_notes":
                i = int(op["index"]) - 1
                prs.slides[i].notes_slide.notes_text_frame.text = str(op.get("notes", ""))
                applied.append(f"set notes on slide {i + 1}")
            elif kind == "set_title":
                i = int(op["index"]) - 1
                slide = prs.slides[i]
                tf = next((s.text_frame for s in slide.shapes
                           if getattr(s, "has_text_frame", False) and s.text_frame.text.strip()), None)
                if tf is None or not tf.paragraphs[0].runs:
                    raise ValueError(f"slide {i + 1} has no text to retitle")
                # the title is the first text box that is not the footer
                for shp in slide.shapes:
                    if getattr(shp, "has_text_frame", False) and shp.text_frame.text.strip() \
                            and shp.top < 1.6 * 914400:
                        shp.text_frame.paragraphs[0].runs[0].text = str(op["title"])
                        break
                else:
                    tf.paragraphs[0].runs[0].text = str(op["title"])
                applied.append(f"retitled slide {i + 1}")
            else:
                warnings.append(f"unknown op {kind!r}")
        except Exception as e:                    # noqa: BLE001
            warnings.append(f"{kind}: {e}")

    if not applied:
        return ToolResult(False, error="nothing was changed: " + "; ".join(warnings or ["no ops"]))
    if in_place:
        journal.record("overwrite", str(src), session=REGISTRY.session_id)
        dest = src
    else:
        dest = _unique(_p(output) if output else src.with_name(src.stem + "-edited.pptx"))
        journal.record("create", str(dest), snapshot=False, session=REGISTRY.session_id)
    try:
        prs.save(str(dest))
        count = len(Presentation(str(dest)).slides)
    except Exception as e:                        # noqa: BLE001
        return ToolResult(False, error=f"could not save: {e}"[:200])
    content = {"path": str(dest), "slides": count, "changes": applied, "verified": True}
    if warnings:
        content["warnings"] = warnings
    return ToolResult(True, content=content,
                      display=f"edited deck: {len(applied)} change(s), {count} slides")


@REGISTRY.tool(
    "presentation_preview",
    "Render a picture of a presentation's first slide (PNG) for a quick visual "
    "check; returns the image path. Pair with look_at_image to review the design.",
    {"type": "object", "properties": {"path": {"type": "string"}}, "required": ["path"]},
    risk=Risk.SAFE, category="read", slow=True)
def presentation_preview(path: str):
    f = _p(path)
    if not f.exists():
        return ToolResult(False, error=f"File not found: {f}")
    png = thumbnail(f)
    if not png:
        return ToolResult(False, error="could not render a preview (Quick Look failed)")
    return ToolResult(True, content={"image": png, "slide": 1},
                      display=f"previewed {f.name}")
