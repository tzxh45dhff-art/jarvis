"""
Word (.docx) output for make_document.

Asked for a Word file, the model used to hand-write python-docx scripts through
run_shell: six steps for a one-page plan (19 Oct 2026), a different look every
time, and no cover, contents, captions or page numbers. This turns the same
HTML that make_document prints to PDF into a native Word document - real
headings (so Word's navigation pane and contents work), real lists and tables,
pictures with captions, diagrams rendered to images, the theme's fonts and
colours - so a Word request is one tool call like a PDF one.
"""
from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import time
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

# Fonts that ship with macOS and Office, closest to each theme's web fonts.
FONTS = {"modern": ("Helvetica Neue", "Helvetica Neue"), "classic": ("Palatino", "Palatino"),
         "minimal": ("Helvetica Neue", "Helvetica Neue"), "warm": ("Avenir Next", "Avenir Next"),
         "forest": ("Avenir Next", "Avenir Next")}
MONO = "Menlo"


def _rgb(hexs: str) -> RGBColor:
    h = hexs.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade(cell_or_par, hexs: str) -> None:
    """Background fill for a table cell or paragraph."""
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexs.lstrip("#"))
    pr.append(shd)


def _border(par, side: str, hexs: str, size: int = 6, space: int = 4) -> None:
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    b = OxmlElement(f"w:{side}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), str(space))
    b.set(qn("w:color"), hexs.lstrip("#"))
    bdr.append(b)


def _field(par, instr: str) -> None:
    """A Word field (PAGE, NUMPAGES) in a paragraph."""
    run = par.add_run()
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
        elif text == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {instr} "
            run._r.append(it)
        else:
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)


def _hyperlink(par, url: str, text: str, color: str) -> None:
    part = par.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                         is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color.lstrip("#"))
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(c)
    rpr.append(u)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    par._p.append(h)


# ------------------------------------------------------------------ diagrams

def render_diagrams(b, theme: str, size: str, build_html, chrome) -> dict[int, Path]:
    """Every Mermaid diagram as a cropped PNG, all drawn in one Chrome pass."""
    if not b.diagrams:
        return {}
    import html as H
    figs = "".join(f"<figure class='diagram' style='break-after: page'><pre class='mermaid' "
                   f"data-index='{i}'>{H.escape(src)}</pre></figure>" for i, src in enumerate(b.diagrams))
    page = b.dir / "diagrams.html"
    page.write_text(build_html("diagrams", "", "", theme, size, figs, [], False, False, ""))
    pdf = b.dir / "diagrams.pdf"
    chrome(["--no-pdf-header-footer", "--virtual-time-budget=15000", f"--print-to-pdf={pdf}",
            page.as_uri()], timeout=120)
    out: dict[int, Path] = {}
    tool = shutil.which("pdftoppm") or "/opt/homebrew/bin/pdftoppm"
    if not pdf.exists() or not Path(tool).exists():
        return out
    shots = Path(tempfile.mkdtemp(prefix="jarvis-dgm-", dir=str(b.dir)))
    subprocess.run([tool, "-r", "200", "-png", str(pdf), str(shots / "d")], capture_output=True, timeout=90)
    from PIL import Image, ImageChops
    for i, s in enumerate(sorted(shots.glob("d*.png"))[:len(b.diagrams)]):
        im = Image.open(s).convert("RGB")
        bg = Image.new("RGB", im.size, (255, 255, 255))
        box = ImageChops.difference(im, bg).getbbox()
        if box:
            pad = 16
            im = im.crop((max(0, box[0] - pad), max(0, box[1] - pad),
                          min(im.width, box[2] + pad), min(im.height, box[3] + pad)))
        p = b.assets / f"diagram-{i + 1}.png"
        im.save(p)
        out[i] = p
    return out


# ------------------------------------------------------------------ writer

class _Writer:
    def __init__(self, doc, th: dict, theme: str, assets: Path, diagrams: dict):
        self.doc, self.th, self.assets, self.diagrams = doc, th, assets, diagrams
        self.body_font, self.head_font = FONTS.get(theme, FONTS["modern"])
        self.fig = 0
        sec = doc.sections[0]
        self.width = sec.page_width - sec.left_margin - sec.right_margin

    # inline -----------------------------------------------------------
    def inline(self, par, node, bold=False, italic=False, code=False):
        for ch in node.children:
            if isinstance(ch, NavigableString):
                text = str(ch)
                if not text:
                    continue
                if not code:
                    text = re.sub(r"\s+", " ", text)
                r = par.add_run(text)
                r.bold, r.italic = bold or None, italic or None
                if code:
                    r.font.name = MONO
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = _rgb(self.th["accent2"])
                continue
            if not isinstance(ch, Tag):
                continue
            n = ch.name
            if n in ("strong", "b"):
                self.inline(par, ch, True, italic, code)
            elif n in ("em", "i"):
                self.inline(par, ch, bold, True, code)
            elif n == "code":
                self.inline(par, ch, bold, italic, True)
            elif n == "br":
                par.add_run().add_break()
            elif n == "a" and ch.get("href", "").startswith(("http://", "https://", "mailto:")):
                _hyperlink(par, ch["href"], ch.get_text(), self.th["accent"])
            else:
                self.inline(par, ch, bold, italic, code)

    # blocks -----------------------------------------------------------
    def block(self, node, level=0):
        if isinstance(node, NavigableString):
            if str(node).strip():
                p = self.doc.add_paragraph()
                p.add_run(str(node).strip())
            return
        if not isinstance(node, Tag):
            return
        n = node.name
        cls = " ".join(node.get("class") or [])
        if n in ("h1", "h2", "h3", "h4", "h5"):
            lvl = min(int(n[1]), 4)
            h = self.doc.add_heading(level=max(1, lvl - 1) if lvl > 1 else 1)
            self.inline(h, node)
            return
        if n == "p":
            if node.find("img") and len(node.get_text(strip=True)) == 0:
                for img in node.find_all("img"):
                    self.picture(img.get("src", ""), img.get("alt", ""), "")
                return
            p = self.doc.add_paragraph()
            self.inline(p, node)
            return
        if n in ("ul", "ol"):
            self.list(node, n == "ol", level)
            return
        if n == "table":
            self.table(node)
            return
        if n == "hr":
            p = self.doc.add_paragraph()
            _border(p, "bottom", self.th["rule"], 8, 1)
            return
        if n == "blockquote":
            for ch in node.children:
                if isinstance(ch, Tag) and ch.name == "p":
                    p = self.doc.add_paragraph(style="Quote")
                    self.inline(p, ch)
                else:
                    self.block(ch, level)
            return
        if n == "pre" or (n == "div" and "code" in cls):
            pre = node if n == "pre" else (node.find("pre") or node)
            p = self.doc.add_paragraph()
            r = p.add_run(pre.get_text().rstrip("\n"))
            r.font.name, r.font.size = MONO, Pt(9)
            _shade(p, self.th["soft"])
            p.paragraph_format.left_indent = Cm(0.3)
            return
        if n == "div" and "callout" in cls:
            self.callout(node)
            return
        if n == "figure":
            self.figure(node)
            return
        for ch in node.children:                    # div, section, main, ...
            self.block(ch, level)

    def list(self, node, ordered: bool, level: int):
        style = ("List Number" if ordered else "List Bullet") + (f" {level + 1}" if level else "")
        for li in node.find_all("li", recursive=False):
            p = self.doc.add_paragraph(style=style)
            nested = []
            for ch in li.children:
                if isinstance(ch, Tag) and ch.name in ("ul", "ol"):
                    nested.append(ch)
                elif isinstance(ch, Tag) and ch.name == "p":
                    self.inline(p, ch)
                else:
                    tmp = BeautifulSoup("<span></span>", "lxml").span
                    tmp.append(ch.__copy__() if isinstance(ch, Tag) else NavigableString(str(ch)))
                    self.inline(p, tmp)
            for sub in nested:
                self.list(sub, sub.name == "ol", min(level + 1, 2))

    def table(self, node):
        rows = node.find_all("tr")
        if not rows:
            return
        ncol = max(len(r.find_all(["td", "th"])) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self._full_width(t, ncol)
        for i, r in enumerate(rows):
            cells = r.find_all(["td", "th"])
            head = i == 0 and (r.find("th") is not None or r.parent.name == "thead")
            for j in range(ncol):
                c = t.cell(i, j)
                c.paragraphs[0].text = ""
                if j < len(cells):
                    self.inline(c.paragraphs[0], cells[j], bold=head)
                if head:
                    _shade(c, self.th["accent"])
                    for run in c.paragraphs[0].runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                elif i % 2 == 0:
                    _shade(c, self.th["soft"])
        self.doc.add_paragraph()

    def _full_width(self, t, ncol: int):
        """Tables span the text column (Word otherwise shrinks them to the
        widest word in each column)."""
        t.autofit = False
        w = int(self.width / max(1, ncol))
        for col in t.columns:
            col.width = Emu(w)
        for row in t.rows:
            for c in row.cells:
                c.width = Emu(w)
        # tblW must sit after tblStyle and before jc/tblLayout/tblLook
        # (schema order); Word offers to "repair" a file that gets it wrong.
        tblPr = t._tbl.tblPr
        for old in tblPr.findall(qn("w:tblW")):
            tblPr.remove(old)
        tw = OxmlElement("w:tblW")
        tw.set(qn("w:w"), "5000")
        tw.set(qn("w:type"), "pct")
        style = tblPr.find(qn("w:tblStyle"))
        if style is not None:
            style.addnext(tw)
        else:
            tblPr.insert(0, tw)

    def callout(self, node):
        title = node.find(class_="callout-title")
        t = self.doc.add_table(rows=1, cols=1)
        self._full_width(t, 1)
        c = t.cell(0, 0)
        _shade(c, self.th["soft"])
        p = c.paragraphs[0]
        if title:
            r = p.add_run(title.get_text(strip=True))
            r.bold = True
            r.font.color.rgb = _rgb(self.th["accent"])
            title.extract()
        first = True
        for ch in node.find_all(["p", "li"]):
            q = p if (first and not title) else c.add_paragraph()
            first = False
            self.inline(q, ch)
        self.doc.add_paragraph()

    def figure(self, node):
        cap_el = node.find("figcaption")
        cap = credit = ""
        if cap_el:
            cr = cap_el.find(class_="credit")
            if cr:
                credit = cr.get_text(" ", strip=True)
                cr.extract()
            lab = cap_el.find(class_="fig-label")
            if lab:
                lab.extract()
            cap = cap_el.get_text(" ", strip=True)
        img = node.find("img")
        pre = node.find("pre", class_="mermaid")
        if pre is not None:
            idx = int(pre.get("data-index", -1))
            path = self.diagrams.get(idx)
            if path:
                self.picture(str(path), cap, credit, absolute=True)
            return
        if img is not None:
            self.picture(img.get("src", ""), cap, credit)

    def picture(self, src: str, cap: str, credit: str, absolute: bool = False):
        p = Path(src) if absolute else self.assets.parent / src
        if not p.exists():
            return
        from PIL import Image
        try:
            w, h = Image.open(p).size
        except Exception:
            return
        maxw = int(self.width * (0.9 if h <= w else 0.6))
        width = Emu(min(maxw, int(Cm(16))))
        par = self.doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(str(p), width=width)
        if cap or credit:
            self.fig += 1
            cp = self.doc.add_paragraph(style="Caption")
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if cap:
                r = cp.add_run(f"Figure {self.fig}  ")
                r.bold = True
                cp.add_run(cap)
            if credit:
                cr = cp.add_run(("  " if cap else "") + credit)
                cr.font.color.rgb = _rgb(self.th["muted"])
                cr.font.size = Pt(8)


def _styles(doc, th: dict, theme: str) -> None:
    body, head = FONTS.get(theme, FONTS["modern"])
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = body
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), body)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _rgb(th["ink"])
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2
    sizes = {1: 18, 2: 14, 3: 12, 4: 11}
    for lvl, sz in sizes.items():
        s = st[f"Heading {lvl}"]
        s.font.name = head
        s.font.size = Pt(sz)
        s.font.bold = True
        s.font.color.rgb = _rgb(th["accent"] if lvl <= 2 else th["ink"])
        s.paragraph_format.space_before = Pt(16 if lvl == 1 else 12)
        s.paragraph_format.space_after = Pt(6)
        rf = s.element.rPr.rFonts
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(k), head)
    cap = st["Caption"]
    cap.font.size = Pt(9)
    cap.font.italic = False
    cap.font.color.rgb = _rgb(th["muted"])
    q = st["Quote"]
    q.font.color.rgb = _rgb(th["muted"])


def write_docx(out: Path, title: str, subtitle: str, author: str, kicker: str, theme: str,
               page_size: str, body_html: str, heads: list, cover: bool, toc: bool,
               assets: Path, diagrams: dict, th: dict) -> dict:
    doc = Document()
    sec = doc.sections[0]
    if page_size == "letter":
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.3))
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2.2)
    _styles(doc, th, theme)
    doc.core_properties.title = title
    if author:
        doc.core_properties.author = author

    if cover:
        for _ in range(6):
            doc.add_paragraph()
        k = doc.add_paragraph()
        r = k.add_run((kicker or "Report").upper())
        r.font.size, r.bold = Pt(10), True
        r.font.color.rgb = _rgb(th["accent2"])
        t = doc.add_paragraph()
        r = t.add_run(title)
        r.font.size, r.bold = Pt(30), True
        r.font.name = FONTS.get(theme, FONTS["modern"])[1]
        r.font.color.rgb = _rgb(th["accent"])
        _border(t, "bottom", th["accent"], 12, 10)
        if subtitle:
            s = doc.add_paragraph()
            r = s.add_run(subtitle)
            r.font.size = Pt(14)
            r.font.color.rgb = _rgb(th["muted"])
        meta = "   ·   ".join(x for x in (author, time.strftime("%d %B %Y").lstrip("0")) if x)
        m = doc.add_paragraph()
        r = m.add_run(meta)
        r.font.size = Pt(10)
        r.font.color.rgb = _rgb(th["muted"])
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    else:
        h = doc.add_heading(title, level=0)
        for r in h.runs:
            r.font.color.rgb = _rgb(th["accent"])
            r.font.name = FONTS.get(theme, FONTS["modern"])[1]
            r.font.bold = True
            r.font.size = Pt(24)

    h2s = [x for x in heads if x[0] == 2]
    if toc and len(h2s) >= 3:
        hd = doc.add_paragraph()
        r = hd.add_run("Contents")
        r.bold, r.font.size = True, Pt(16)
        r.font.color.rgb = _rgb(th["accent"])
        for i, (_, name, _) in enumerate(h2s, 1):
            p = doc.add_paragraph()
            r = p.add_run(f"{i:02d}   ")
            r.bold = True
            r.font.color.rgb = _rgb(th["accent2"])
            p.add_run(name)
            _border(p, "bottom", th["rule"], 4, 2)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # page footer: title on the left, page x of y on the right
    foot = sec.footer.paragraphs[0]
    foot.text = ""
    r = foot.add_run(title[:80] + "\t\t")
    r.font.size = Pt(8)
    r.font.color.rgb = _rgb(th["muted"])
    _field(foot, "PAGE")
    foot.add_run(" / ").font.size = Pt(8)
    _field(foot, "NUMPAGES")
    for run in foot.runs:
        run.font.size = Pt(8)
    _border(foot, "top", th["rule"], 4, 4)
    sec.different_first_page_header_footer = bool(cover)

    w = _Writer(doc, th, theme, assets, diagrams)
    soup = BeautifulSoup(body_html, "lxml")
    root = soup.body or soup
    for ch in list(root.children):
        w.block(ch)
    doc.save(str(out))
    return {"figures": w.fig}


def preview(docx: Path, folder: Path) -> str | None:
    """A first-page picture via Quick Look, to look at before reporting."""
    folder.mkdir(parents=True, exist_ok=True)
    try:
        subprocess.run(["qlmanage", "-t", "-s", "1200", "-o", str(folder), str(docx)],
                       capture_output=True, timeout=40)
    except Exception:
        return None
    p = folder / (docx.name + ".png")
    return str(p) if p.exists() else None
