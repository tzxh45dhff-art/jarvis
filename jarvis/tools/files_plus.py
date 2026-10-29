"""
Documents and files beyond plain text.

`read_file` only ever read text, so a PDF, a Word file, a spreadsheet, a slide
deck or a photo of a receipt was a closed book. These tools open all of them,
and turn one kind of file into another:

    read_document  PDF (scanned too, via OCR), DOCX, XLSX/CSV, PPTX, RTF, HTML,
                   images (OCR), archives and media (metadata)
    pdf_tool       info, merge, extract pages, split, rotate, to images
    convert_file   images, audio, video, documents, csv<->xlsx
    archive_tool   zip, unzip, list

Everything that writes produces a *new* file next to the source (never
overwrites), checks it exists and is non-empty, and is journalled like every
other file JARVIS creates.
"""
from __future__ import annotations

import csv
import io
import logging
import os
import re
import shutil
import subprocess
import tarfile
import time
import zipfile
from pathlib import Path

from ..safety import journal
from ..safety.policy import Risk, Verdict, classify_path_write
from .fs import _p
from .registry import REGISTRY, ToolResult

log = logging.getLogger("jarvis.files")

PAGE = 9000
_BIN = ("/opt/homebrew/bin", "/usr/local/bin", "/usr/bin", "/bin")


def _which(name: str) -> str | None:
    for d in _BIN:
        c = os.path.join(d, name)
        if os.path.exists(c) and os.access(c, os.X_OK):
            return c
    return shutil.which(name)


def _run(cmd: list[str], timeout: int = 120) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, text=True, timeout=timeout,
                          errors="replace")


def _unique(path: Path) -> Path:
    if not path.exists():
        return path
    for i in range(2, 1000):
        c = path.with_name(f"{path.stem} ({i}){path.suffix}")
        if not c.exists():
            return c
    return path.with_name(f"{path.stem}-{int(time.time())}{path.suffix}")


def _parse_pages(spec: str, total: int) -> list[int]:
    """'1-3,7' -> [0, 1, 2, 6] (zero-based, in range, in order, no repeats)."""
    if not spec or str(spec).strip().lower() in ("all", "*"):
        return list(range(total))
    out: list[int] = []
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, _, b = part.partition("-")
            lo = int(a) if a else 1
            hi = int(b) if b else total
        else:
            lo = hi = int(part)
        for n in range(max(1, lo), min(total, hi) + 1):
            if n - 1 not in out:
                out.append(n - 1)
    return out


# ---------------------------------------------------------------- OCR

def ocr_image(path: str) -> str:
    """Text in an image file, top to bottom, using Apple's Vision framework
    (falls back to tesseract). Empty string if nothing readable."""
    try:
        import Vision
        from Foundation import NSURL
        lines: list[tuple[float, float, str]] = []

        def handler(req, err):
            for obs in req.results() or []:
                cand = obs.topCandidates_(1)
                if cand:
                    bb = obs.boundingBox()
                    lines.append((1.0 - bb.origin.y - bb.size.height, bb.origin.x,
                                  str(cand[0].string())))

        req = Vision.VNRecognizeTextRequest.alloc().initWithCompletionHandler_(handler)
        req.setRecognitionLevel_(0)               # accurate
        req.setUsesLanguageCorrection_(True)
        h = Vision.VNImageRequestHandler.alloc().initWithURL_options_(
            NSURL.fileURLWithPath_(str(path)), None)
        ok, _ = h.performRequests_error_([req], None)
        if ok and lines:
            # group into rows: same line if their tops are within ~1.2% of height
            lines.sort()
            rows, cur, last = [], [], None
            for top, left, text in lines:
                if last is not None and abs(top - last) > 0.012:
                    rows.append(sorted(cur, key=lambda t: t[0]))
                    cur = []
                cur.append((left, text))
                last = top
            if cur:
                rows.append(sorted(cur, key=lambda t: t[0]))
            return "\n".join("  ".join(t for _, t in r) for r in rows)
    except Exception as e:                       # noqa: BLE001
        log.info("Vision OCR unavailable: %s", e)
    tess = _which("tesseract")
    if tess:
        r = _run([tess, str(path), "stdout", "-l", "eng", "--psm", "6"], 90)
        return r.stdout.strip()
    return ""


# ---------------------------------------------------------------- readers

def _read_pdf(f: Path, pages_spec: str) -> dict:
    try:
        from pypdf import PdfReader
        rd = PdfReader(str(f))
        if rd.is_encrypted:
            try:
                rd.decrypt("")
            except Exception:
                return {"error": "this PDF is password-protected"}
        total = len(rd.pages)
        meta = rd.metadata
    except Exception as e:
        return {"error": f"could not open PDF: {e}"}
    sel = _parse_pages(pages_spec, total)
    if not sel:
        return {"error": f"no such pages (the document has {total})"}
    text = ""
    tool = _which("pdftotext")
    if tool:
        first, last = sel[0] + 1, sel[-1] + 1
        r = _run([tool, "-layout", "-f", str(first), "-l", str(last), str(f), "-"], 90)
        text = r.stdout
    if not text.strip():
        text = "\n\n".join(
            f"--- page {i + 1} ---\n{(rd.pages[i].extract_text() or '').strip()}" for i in sel)
    out: dict = {"kind": "pdf", "pages_total": total,
                 "title": (getattr(meta, "title", "") or "") if meta else ""}
    per_page = len(text.strip()) / max(1, len(sel))
    if per_page < 40 and _which("pdftoppm"):
        # Scanned: there is no text layer, so read the page images.
        import tempfile
        parts = []
        with tempfile.TemporaryDirectory() as td:
            for i in sel[:8]:
                pre = os.path.join(td, f"p{i + 1}")
                _run([_which("pdftoppm"), "-r", "170", "-f", str(i + 1), "-l", str(i + 1),
                      "-png", str(f), pre], 60)
                for img in sorted(Path(td).glob(f"p{i + 1}*.png")):
                    parts.append(f"--- page {i + 1} ---\n{ocr_image(str(img))}")
        text = "\n\n".join(parts)
        out["ocr"] = True
        if len(sel) > 8:
            out["note"] = "scanned PDF: only the first 8 requested pages were read"
    out["text"] = text
    return out


def _read_docx(f: Path) -> dict:
    import docx
    d = docx.Document(str(f))
    parts = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "") if p.style is not None else ""
        parts.append(("# " + t) if style.startswith("Heading") or style == "Title" else t)
    for ti, tbl in enumerate(d.tables, 1):
        parts.append(f"\n[table {ti}]")
        for row in tbl.rows:
            parts.append(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
    return {"kind": "docx", "text": "\n".join(parts)}


def _read_xlsx(f: Path, max_rows: int = 60) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(f), read_only=True, data_only=True)
    parts, sheets = [], []
    for ws in wb.worksheets:
        sheets.append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column})
        parts.append(f"## sheet: {ws.title}  ({ws.max_row} rows x {ws.max_column} columns)")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                parts.append(f"... {max(0, (ws.max_row or 0) - max_rows)} more rows")
                break
            parts.append("\t".join("" if c is None else str(c) for c in row))
    return {"kind": "xlsx", "sheets": sheets, "text": "\n".join(parts)}


def _read_csv(f: Path, max_rows: int = 80) -> dict:
    raw = f.read_bytes()[:3_000_000].decode("utf-8-sig", "replace")
    dialect = csv.excel_tab if f.suffix.lower() == ".tsv" else csv.excel
    try:
        dialect = csv.Sniffer().sniff(raw[:4000], delimiters=",;\t|")
    except Exception:
        pass
    rows = list(csv.reader(io.StringIO(raw), dialect))
    head = rows[:max_rows]
    text = "\n".join("\t".join(r) for r in head)
    if len(rows) > max_rows:
        text += f"\n... {len(rows) - max_rows} more rows (file has {len(rows):,}+ rows)"
    return {"kind": "csv", "columns": rows[0] if rows else [], "text": text}


def _read_pptx(f: Path) -> dict:
    from pptx import Presentation
    prs = Presentation(str(f))
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"--- slide {i} ---")
        for shp in slide.shapes:
            if getattr(shp, "has_text_frame", False) and shp.text_frame.text.strip():
                parts.append(shp.text_frame.text.strip())
            if getattr(shp, "has_table", False) and shp.has_table:
                for row in shp.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("[notes] " + slide.notes_slide.notes_text_frame.text.strip())
    return {"kind": "pptx", "slides": len(prs.slides), "text": "\n".join(parts)}


def _read_textutil(f: Path) -> dict:
    r = _run(["/usr/bin/textutil", "-convert", "txt", "-stdout", str(f)], 60)
    if r.returncode != 0:
        return {"error": f"could not convert: {r.stderr.strip()[:160]}"}
    return {"kind": f.suffix.lstrip(".").lower() or "document", "text": r.stdout}


def _read_media(f: Path) -> dict:
    probe = _which("ffprobe")
    if not probe:
        return {"kind": "media", "text": f"{f.name}: media file ({f.stat().st_size:,} bytes)"}
    r = _run([probe, "-v", "quiet", "-print_format", "json", "-show_format",
              "-show_streams", str(f)], 30)
    import json
    try:
        j = json.loads(r.stdout)
    except Exception:
        return {"error": "could not read media metadata"}
    fmt = j.get("format", {})
    streams = [{"type": s.get("codec_type"), "codec": s.get("codec_name"),
                "size": f"{s.get('width')}x{s.get('height')}" if s.get("width") else None,
                "channels": s.get("channels")} for s in j.get("streams", [])]
    return {"kind": "media", "duration_seconds": round(float(fmt.get("duration", 0) or 0), 1),
            "format": fmt.get("format_long_name"), "streams": streams,
            "text": f"{f.name}: {round(float(fmt.get('duration', 0) or 0), 1)} s, "
                    + ", ".join(f"{s['type']}:{s['codec']}" for s in streams)}


_IMAGE = {".png", ".jpg", ".jpeg", ".heic", ".heif", ".tif", ".tiff", ".gif", ".bmp", ".webp"}
_MEDIA = {".mp3", ".m4a", ".wav", ".flac", ".aac", ".ogg", ".mp4", ".mov", ".mkv",
          ".avi", ".webm", ".m4v"}


@REGISTRY.tool(
    "read_document",
    "Read the text of a document of any common kind: PDF (including scanned "
    "pages, via OCR), Word (.docx), Excel (.xlsx) and CSV, PowerPoint (.pptx), "
    "RTF/Pages-exported text, HTML, and images (OCR). Use this instead of "
    "read_file for anything that is not plain text. For PDFs pass pages like "
    "'1-3' to read part of a long one. Long documents come in parts: call "
    "again with the given offset.",
    {"type": "object", "properties": {
        "path": {"type": "string"},
        "pages": {"type": "string", "description": "PDF page range, e.g. '1-5,9'"},
        "offset": {"type": "integer", "description": "Character to start from"},
        "max_chars": {"type": "integer", "description": "Default 9000"}},
     "required": ["path"]},
    risk=Risk.SAFE, category="read", parallel_safe=True, slow=True)
def read_document(path: str, pages: str = "", offset: int = 0, max_chars: int = PAGE):
    f = _p(path)
    if not f.exists():
        return ToolResult(False, error=f"File not found: {f}")
    if f.is_dir():
        return ToolResult(False, error=f"{f} is a folder; use list_dir")
    ext = f.suffix.lower()
    try:
        if ext == ".pdf":
            d = _read_pdf(f, pages)
        elif ext in (".docx", ".docm"):
            d = _read_docx(f)
        elif ext in (".xlsx", ".xlsm"):
            d = _read_xlsx(f)
        elif ext in (".csv", ".tsv"):
            d = _read_csv(f)
        elif ext == ".pptx":
            d = _read_pptx(f)
        elif ext in (".rtf", ".rtfd", ".doc", ".odt", ".html", ".htm", ".webarchive", ".wordml"):
            d = _read_textutil(f)
        elif ext in _IMAGE:
            txt = ocr_image(str(f))
            d = {"kind": "image", "ocr": True, "text": txt,
                 "note": ("text recognised in the image; use look_at_image to "
                          "understand what it shows") if txt else
                         "no readable text found; use look_at_image to see what it shows"}
        elif ext in _MEDIA:
            d = _read_media(f)
        elif ext in (".zip",):
            with zipfile.ZipFile(f) as z:
                names = z.namelist()
            d = {"kind": "zip", "text": f"{len(names)} entries:\n" + "\n".join(names[:200])}
        else:
            raw = f.read_bytes()[:1_500_000]
            if b"\x00" in raw[:4096]:
                return ToolResult(False, error=f"{f.name} looks like a binary file "
                                               f"({ext or 'no extension'}); not a document I can read")
            d = {"kind": "text", "text": raw.decode("utf-8", "replace")}
    except ImportError as e:
        return ToolResult(False, error=f"missing reader for {ext}: {e}")
    except Exception as e:                       # noqa: BLE001
        return ToolResult(False, error=f"could not read {f.name}: {e}"[:300])

    if d.get("error"):
        return ToolResult(False, error=d["error"])
    full = d.pop("text", "") or ""
    size = max(500, min(int(max_chars or PAGE), 20000))
    start = max(0, int(offset or 0))
    part = full[start:start + size]
    d.update({"path": str(f), "total_chars": len(full), "offset": start, "content": part})
    if start + size < len(full):
        d["more"] = True
        d["next_offset"] = start + size
        d["note"] = (d.get("note", "") + f" showing characters {start:,}-"
                     f"{start + len(part):,} of {len(full):,}; call again with "
                     f"offset={start + size} for the rest").strip()
    if not full.strip() and not d.get("note"):
        d["warning"] = "no text was found in this file"
    return ToolResult(True, content=d,
                      display=f"read {f.name} ({d.get('kind')}, {len(full):,} chars)")


# ---------------------------------------------------------------- pdf tool

def _assess_out(args: dict) -> Verdict:
    target = args.get("output") or args.get("path") or ""
    if not target and args.get("paths"):
        target = args["paths"][0]
    if target:
        v = classify_path_write(str(target))
        if v.risk >= Risk.FORBIDDEN:
            return v
    return Verdict(Risk.LOW, "creates a new file", category="write_new")


def _finish(target: Path, session: str = "") -> dict:
    if not target.exists() or target.stat().st_size == 0:
        raise RuntimeError("the output file was not created")
    journal.record("create", str(target), snapshot=False,
                   session=session or REGISTRY.session_id)
    return {"path": str(target), "bytes": target.stat().st_size, "verified": True}


@REGISTRY.tool(
    "pdf_tool",
    "Work with PDF files. action: info (page count, title), merge (paths -> one "
    "PDF), extract (pages -> new PDF), split (one file per page), rotate "
    "(pages by angle 90/180/270), to_images (PNG per page). Never overwrites: "
    "new files are named next to the source.",
    {"type": "object", "properties": {
        "action": {"type": "string",
                   "enum": ["info", "merge", "extract", "split", "rotate", "to_images"]},
        "path": {"type": "string", "description": "Source PDF (all actions but merge)"},
        "paths": {"type": "array", "items": {"type": "string"},
                  "description": "PDFs to merge, in order"},
        "pages": {"type": "string", "description": "e.g. '1-3,7'; default all"},
        "angle": {"type": "integer"},
        "output": {"type": "string", "description": "Optional output file or folder"}},
     "required": ["action"]},
    risk=Risk.LOW, category="write_new", parallel_safe=False, slow=True,
    assess=_assess_out)
def pdf_tool(action: str, path: str = "", paths: list | None = None, pages: str = "",
             angle: int = 90, output: str = ""):
    try:
        from pypdf import PdfReader, PdfWriter
    except ImportError:
        return ToolResult(False, error="pypdf is not installed")
    action = (action or "").lower()

    def reader(p: Path):
        rd = PdfReader(str(p))
        if rd.is_encrypted:
            rd.decrypt("")
        return rd

    try:
        if action == "merge":
            srcs = [_p(x) for x in (paths or [])]
            if len(srcs) < 2:
                return ToolResult(False, error="give at least two PDFs in 'paths'")
            for s in srcs:
                if not s.exists():
                    return ToolResult(False, error=f"File not found: {s}")
            w = PdfWriter()
            for s in srcs:
                for pg in reader(s).pages:
                    w.add_page(pg)
            out = _unique(_p(output) if output else srcs[0].with_name(srcs[0].stem + "-merged.pdf"))
            with open(out, "wb") as fh:
                w.write(fh)
            r = _finish(out)
            r["pages"] = len(w.pages)
            return ToolResult(True, content=r, display=f"merged {len(srcs)} PDFs ({len(w.pages)} pages)")

        f = _p(path)
        if not f.exists():
            return ToolResult(False, error=f"File not found: {f}")
        rd = reader(f)
        total = len(rd.pages)
        if action == "info":
            m = rd.metadata
            return ToolResult(True, content={
                "path": str(f), "pages": total, "title": (getattr(m, "title", "") or "") if m else "",
                "author": (getattr(m, "author", "") or "") if m else "",
                "bytes": f.stat().st_size, "encrypted": bool(rd.is_encrypted)},
                display=f"{f.name}: {total} pages")
        sel = _parse_pages(pages, total)
        if not sel:
            return ToolResult(False, error=f"no such pages (the document has {total})")
        if action == "extract":
            w = PdfWriter()
            for i in sel:
                w.add_page(rd.pages[i])
            label = pages.replace(",", "_").replace(" ", "") or "all"
            out = _unique(_p(output) if output else f.with_name(f"{f.stem}-pages-{label}.pdf"))
            with open(out, "wb") as fh:
                w.write(fh)
            r = _finish(out)
            r["pages"] = len(sel)
            return ToolResult(True, content=r, display=f"extracted {len(sel)} pages")
        if action == "rotate":
            if angle % 90:
                return ToolResult(False, error="angle must be 90, 180 or 270")
            w = PdfWriter()
            for i, pg in enumerate(rd.pages):
                if i in sel:
                    pg.rotate(angle)
                w.add_page(pg)
            out = _unique(_p(output) if output else f.with_name(f"{f.stem}-rotated.pdf"))
            with open(out, "wb") as fh:
                w.write(fh)
            r = _finish(out)
            r["rotated_pages"] = len(sel)
            return ToolResult(True, content=r, display=f"rotated {len(sel)} pages by {angle}")
        if action == "split":
            if len(sel) > 200:
                return ToolResult(False, error="refusing to split into more than 200 files")
            folder = _unique(_p(output) if output else f.with_name(f.stem + "-pages"))
            folder.mkdir(parents=True, exist_ok=True)
            made = []
            for i in sel:
                w = PdfWriter()
                w.add_page(rd.pages[i])
                target = folder / f"{f.stem}-{i + 1:03d}.pdf"
                with open(target, "wb") as fh:
                    w.write(fh)
                made.append(target.name)
            journal.record("create", str(folder), snapshot=False, session=REGISTRY.session_id)
            return ToolResult(True, content={"folder": str(folder), "files": made[:50],
                                             "count": len(made), "verified": True},
                              display=f"split into {len(made)} files")
        if action == "to_images":
            tool = _which("pdftoppm")
            if not tool:
                return ToolResult(False, error="pdftoppm is not installed (brew install poppler)")
            if len(sel) > 40:
                return ToolResult(False, error="refusing to render more than 40 pages")
            folder = _unique(_p(output) if output else f.with_name(f.stem + "-images"))
            folder.mkdir(parents=True, exist_ok=True)
            for i in sel:
                _run([tool, "-r", "150", "-f", str(i + 1), "-l", str(i + 1), "-png",
                      str(f), str(folder / f"{f.stem}-{i + 1:03d}")], 90)
            made = sorted(p.name for p in folder.glob("*.png"))
            if not made:
                return ToolResult(False, error="no images were produced")
            journal.record("create", str(folder), snapshot=False, session=REGISTRY.session_id)
            return ToolResult(True, content={"folder": str(folder), "files": made[:50],
                                             "count": len(made), "verified": True},
                              display=f"rendered {len(made)} page images")
        return ToolResult(False, error=f"unknown action {action!r}")
    except Exception as e:                       # noqa: BLE001
        return ToolResult(False, error=f"PDF operation failed: {e}"[:300])


# ---------------------------------------------------------------- convert

_SIPS_FORMATS = {"png": "png", "jpg": "jpeg", "jpeg": "jpeg", "heic": "heic", "tiff": "tiff",
                 "tif": "tiff", "gif": "gif", "bmp": "bmp", "pdf": "pdf"}
_AUDIO_OUT = {"mp3", "m4a", "wav", "flac", "aac", "ogg", "opus"}
_VIDEO_OUT = {"mp4", "mov", "mkv", "webm", "gif", "m4v"}
_TEXTUTIL_OUT = {"txt": "txt", "rtf": "rtf", "html": "html", "docx": "docx",
                 "doc": "doc", "odt": "odt", "webarchive": "webarchive"}


def _convert_image(src: Path, to: str, out: Path, max_dim: int, quality: int):
    if to == "webp" or src.suffix.lower() == ".webp":
        from PIL import Image
        im = Image.open(src)
        if max_dim:
            im.thumbnail((max_dim, max_dim))
        if to in ("jpg", "jpeg"):
            im = im.convert("RGB")
        im.save(out, quality=quality or 85)
        return
    cmd = ["/usr/bin/sips", "-s", "format", _SIPS_FORMATS[to]]
    if quality and _SIPS_FORMATS[to] in ("jpeg", "heic"):
        cmd += ["-s", "formatOptions", str(max(1, min(100, quality)))]
    if max_dim:
        cmd += ["--resampleHeightWidthMax", str(max_dim)]
    cmd += [str(src), "--out", str(out)]
    r = _run(cmd, 120)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.strip()[:200] or "sips failed")


def _convert_av(src: Path, to: str, out: Path, start: str, end: str, audio_only: bool):
    ff = _which("ffmpeg")
    if not ff:
        raise RuntimeError("ffmpeg is not installed (brew install ffmpeg)")
    cmd = [ff, "-y", "-v", "error", "-i", str(src)]
    if start:
        cmd = [ff, "-y", "-v", "error", "-ss", str(start), "-i", str(src)]
    if end:
        cmd += ["-to", str(end)] if not start else ["-t", str(_dur(start, end))]
    if to == "gif":
        cmd += ["-vf", "fps=12,scale=480:-1:flags=lanczos", "-loop", "0"]
    elif to in _AUDIO_OUT or audio_only:
        cmd += ["-vn"]
        if to == "mp3":
            cmd += ["-codec:a", "libmp3lame", "-q:a", "2"]
        elif to in ("m4a", "aac"):
            cmd += ["-codec:a", "aac", "-b:a", "192k"]
    elif to in ("mp4", "mov", "m4v"):
        cmd += ["-codec:v", "libx264", "-preset", "fast", "-crf", "23",
                "-codec:a", "aac", "-movflags", "+faststart", "-pix_fmt", "yuv420p"]
    cmd += [str(out)]
    r = _run(cmd, 900)
    if r.returncode != 0:
        raise RuntimeError(r.stderr.strip()[-240:] or "ffmpeg failed")


def _secs(t: str) -> float:
    parts = [float(x) for x in str(t).split(":")]
    s = 0.0
    for p in parts:
        s = s * 60 + p
    return s


def _dur(start: str, end: str) -> float:
    return max(0.1, _secs(end) - _secs(start))


@REGISTRY.tool(
    "convert_file",
    "Convert a file to another format and save it next to the original (never "
    "overwrites). Images: png/jpg/heic/tiff/gif/webp/pdf, with optional "
    "max_dimension (resize) and quality. Audio/video via ffmpeg: mp3, m4a, wav, "
    "mp4, mov, gif, webm... with optional start/end ('1:30') to trim, "
    "audio_only to extract sound. Documents: txt/rtf/html/docx/odt, markdown to "
    "html or pdf, csv<->xlsx, json to csv.",
    {"type": "object", "properties": {
        "path": {"type": "string"},
        "to": {"type": "string", "description": "Target format, e.g. 'pdf', 'mp3', 'png'"},
        "output": {"type": "string", "description": "Optional output file path"},
        "max_dimension": {"type": "integer", "description": "Resize images so the long side is this many pixels"},
        "quality": {"type": "integer", "description": "1-100 for jpg/heic/webp"},
        "start": {"type": "string"}, "end": {"type": "string"},
        "audio_only": {"type": "boolean"}},
     "required": ["path", "to"]},
    risk=Risk.LOW, category="write_new", parallel_safe=False, slow=True,
    assess=_assess_out)
def convert_file(path: str, to: str, output: str = "", max_dimension: int = 0,
                 quality: int = 0, start: str = "", end: str = "", audio_only: bool = False):
    src = _p(path)
    if not src.exists():
        return ToolResult(False, error=f"File not found: {src}")
    to = (to or "").lower().lstrip(".")
    if to == "jpeg":
        to = "jpg"
    ext = src.suffix.lower().lstrip(".")
    if not to:
        return ToolResult(False, error="say which format to convert to")
    if ext == to:
        return ToolResult(False, error=f"{src.name} is already {to}")
    out = _unique(_p(output) if output else src.with_suffix("." + to))
    t0 = time.time()
    try:
        if ("." + ext) in _IMAGE and to in set(_SIPS_FORMATS) | {"webp"}:
            _convert_image(src, to, out, int(max_dimension or 0), int(quality or 0))
        elif (("." + ext) in _MEDIA or ext in ("gif",)) and (to in _AUDIO_OUT or to in _VIDEO_OUT):
            _convert_av(src, to, out, start, end, audio_only)
        elif ext in ("md", "markdown", "txt") and to == "html":
            import markdown
            body = markdown.markdown(src.read_text(errors="replace"),
                                     extensions=["tables", "fenced_code"])
            out.write_text("<!doctype html><meta charset=utf-8><title>"
                           + src.stem + "</title><body style='font-family:-apple-system;"
                           "max-width:760px;margin:2em auto;line-height:1.5'>" + body + "</body>")
        elif ext in ("md", "markdown", "txt") and to == "pdf":
            from .docs import _pdf
            ok, msg = _pdf(out, src.stem, src.read_text(errors="replace"))
            if not ok:
                raise RuntimeError(msg)
        elif ext in ("rtf", "doc", "docx", "odt", "html", "htm", "txt", "webarchive") \
                and to in _TEXTUTIL_OUT:
            r = _run(["/usr/bin/textutil", "-convert", _TEXTUTIL_OUT[to], str(src),
                      "-output", str(out)], 90)
            if r.returncode != 0:
                raise RuntimeError(r.stderr.strip()[:200])
        elif ext == "csv" and to == "xlsx":
            import pandas as pd
            pd.read_csv(src).to_excel(out, index=False)
        elif ext in ("xlsx", "xlsm") and to == "csv":
            import pandas as pd
            pd.read_excel(src).to_csv(out, index=False)
        elif ext == "json" and to == "csv":
            import pandas as pd
            pd.read_json(src).to_csv(out, index=False)
        elif ext == "csv" and to == "json":
            import pandas as pd
            pd.read_csv(src).to_json(out, orient="records", indent=2)
        else:
            return ToolResult(False, error=f"don't know how to convert .{ext} to .{to}")
        r = _finish(out)
    except Exception as e:                       # noqa: BLE001
        out.unlink(missing_ok=True) if out.exists() and out.stat().st_size == 0 else None
        return ToolResult(False, error=f"conversion failed: {e}"[:300])
    r.update({"from": ext, "to": to, "seconds": round(time.time() - t0, 1)})
    return ToolResult(True, content=r,
                      display=f"converted {src.name} -> {Path(r['path']).name} "
                              f"({r['bytes'] / 1e6:.1f} MB)" if r["bytes"] > 1e5
                      else f"converted {src.name} -> {Path(r['path']).name}")


# ---------------------------------------------------------------- archives

MAX_UNZIP = 4 * 1024 ** 3


@REGISTRY.tool(
    "archive_tool",
    "Create or open archives. action: zip (paths -> .zip), unzip (extract a "
    ".zip/.tar/.tar.gz/.tgz into a new folder), list (show contents without "
    "extracting). Never overwrites.",
    {"type": "object", "properties": {
        "action": {"type": "string", "enum": ["zip", "unzip", "list"]},
        "path": {"type": "string", "description": "Archive to unzip or list"},
        "paths": {"type": "array", "items": {"type": "string"},
                  "description": "Files or folders to zip"},
        "output": {"type": "string"}},
     "required": ["action"]},
    risk=Risk.LOW, category="write_new", parallel_safe=False, slow=True,
    assess=_assess_out)
def archive_tool(action: str, path: str = "", paths: list | None = None, output: str = ""):
    action = (action or "").lower()
    try:
        if action == "zip":
            srcs = [_p(x) for x in (paths or ([path] if path else []))]
            if not srcs:
                return ToolResult(False, error="give the files or folders to zip in 'paths'")
            for s in srcs:
                if not s.exists():
                    return ToolResult(False, error=f"File not found: {s}")
            name = (srcs[0].stem if len(srcs) == 1 else "archive") + ".zip"
            out = _unique(_p(output) if output else srcs[0].with_name(name))
            n = 0
            with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
                for s in srcs:
                    if s.is_dir():
                        for root, dirs, files in os.walk(s):
                            dirs[:] = [d for d in dirs if d not in (".git", "node_modules", "__pycache__")]
                            for fn in files:
                                fp = Path(root) / fn
                                z.write(fp, str(Path(s.name) / fp.relative_to(s)))
                                n += 1
                    else:
                        z.write(s, s.name)
                        n += 1
            r = _finish(out)
            r["files"] = n
            return ToolResult(True, content=r, display=f"zipped {n} files -> {out.name}")

        f = _p(path)
        if not f.exists():
            return ToolResult(False, error=f"File not found: {f}")
        is_zip = zipfile.is_zipfile(f)
        is_tar = tarfile.is_tarfile(f)
        if not (is_zip or is_tar):
            return ToolResult(False, error=f"{f.name} is not a zip or tar archive")
        if action == "list":
            if is_zip:
                with zipfile.ZipFile(f) as z:
                    infos = z.infolist()
                    names = [f"{i.filename}  ({i.file_size:,} bytes)" for i in infos[:300]]
                    total = sum(i.file_size for i in infos)
                    count = len(infos)
            else:
                with tarfile.open(f) as t:
                    ms = t.getmembers()
                    names = [f"{m.name}  ({m.size:,} bytes)" for m in ms[:300]]
                    total = sum(m.size for m in ms)
                    count = len(ms)
            return ToolResult(True, content={"path": str(f), "entries": count,
                                             "uncompressed_bytes": total, "files": names},
                              display=f"{f.name}: {count} entries")
        if action == "unzip":
            stem = f.name
            for suf in (".tar.gz", ".tgz", ".tar", ".zip"):
                if stem.endswith(suf):
                    stem = stem[: -len(suf)]
                    break
            dest = _unique(_p(output) if output else f.with_name(stem))
            dest.mkdir(parents=True, exist_ok=True)
            root = dest.resolve()

            def safe(name: str) -> Path:
                tgt = (dest / name).resolve()
                if root != tgt and root not in tgt.parents:
                    raise RuntimeError(f"unsafe path in archive: {name}")
                return tgt

            if is_zip:
                with zipfile.ZipFile(f) as z:
                    if sum(i.file_size for i in z.infolist()) > MAX_UNZIP:
                        raise RuntimeError("archive expands to more than 4 GB; refusing")
                    for i in z.infolist():
                        safe(i.filename)
                    z.extractall(dest)
                    count = len(z.infolist())
            else:
                with tarfile.open(f) as t:
                    ms = t.getmembers()
                    if sum(m.size for m in ms) > MAX_UNZIP:
                        raise RuntimeError("archive expands to more than 4 GB; refusing")
                    for m in ms:
                        safe(m.name)
                        if m.issym() or m.islnk():
                            raise RuntimeError(f"archive contains a link: {m.name}")
                    t.extractall(dest)
                    count = len(ms)
            journal.record("create", str(dest), snapshot=False, session=REGISTRY.session_id)
            return ToolResult(True, content={"folder": str(dest), "entries": count,
                                             "verified": dest.exists()},
                              display=f"extracted {count} entries -> {dest.name}")
        return ToolResult(False, error=f"unknown action {action!r}")
    except Exception as e:                       # noqa: BLE001
        return ToolResult(False, error=f"archive operation failed: {e}"[:300])
